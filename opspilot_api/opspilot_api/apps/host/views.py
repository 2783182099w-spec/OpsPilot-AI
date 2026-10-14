from io import BytesIO
from django.db.models import Q

from django.conf import settings
from django.shortcuts import get_object_or_404
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.generics import ListCreateAPIView
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework.viewsets import ModelViewSet

from .models import Host, HostCategory
from .permissions import accessible_hosts
from .serializers import HostCategoryModelSerializer, HostModelSerializer
from opspilot_api.utils.ssh import SSHParamiko


class HostCategoryListAPIView(ListCreateAPIView):
    """主机类别：列表 / 新建"""

    serializer_class = HostCategoryModelSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return HostCategory.objects.filter(is_show=True, is_deleted=False).order_by("orders", "-id")


class HostModelViewSet(ModelViewSet):
    """主机信息：增删改查"""

    serializer_class = HostModelSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = accessible_hosts(self.request.user).select_related("category").order_by("-id")
        category_id = self.request.query_params.get("category")
        if category_id:
            queryset = queryset.filter(category_id=category_id)
        search = self.request.query_params.get("search", "")[:100]
        if search:
            queryset = queryset.filter(Q(name__icontains=search) | Q(ip_addr__icontains=search) | Q(description__icontains=search))
        return queryset

    def list(self, request, *args, **kwargs):
        if "page" not in request.query_params:
            return super().list(request, *args, **kwargs)
        from rest_framework.pagination import PageNumberPagination
        paginator = PageNumberPagination()
        paginator.page_size = 20
        page = paginator.paginate_queryset(self.get_queryset(), request)
        return paginator.get_paginated_response(self.get_serializer(page, many=True).data)


class HostKeyUploadView(APIView):
    """用登录密码连接一次目标主机，把系统全局 SSH 公钥写入 authorized_keys。"""

    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        password = request.data.get("password")
        if not password:
            return Response({"error": "请输入主机登录密码"}, status=status.HTTP_400_BAD_REQUEST)

        host = get_object_or_404(accessible_hosts(request.user), pk=pk)
        _, public_key = HostModelSerializer.ensure_global_key()

        ssh = SSHParamiko(host.ip_addr, host.port, host.username, password=str(password))
        if ssh.get_connected_client() is None:
            return Response(
                {"error": f"无法使用该密码连接到主机：{ssh.error_message}"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            ssh.upload_key(public_key)
        except Exception as exc:  # noqa: BLE001
            return Response(
                {"error": f"公钥写入失败：{exc}"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        finally:
            ssh.close()

        return Response({"message": "公钥已同步到目标主机"})


def _format_errors(errors):
    """把 DRF 的错误字典压成一行可读文本。"""
    parts = []
    for field, messages in errors.items():
        text = "；".join(str(message) for message in messages)
        parts.append(text if field == "non_field_errors" else f"{field}: {text}")
    return "；".join(parts)


class HostExcelView(APIView):
    """Excel 批量导入主机。

    模板约定：第二张工作表为数据表，列顺序为
    类别名称、主机名称、连接地址、端口、登录用户、连接密码、备注。
    """

    parser_classes = [MultiPartParser, FormParser]
    permission_classes = [IsAuthenticated]

    def post(self, request):
        host_excel = request.FILES.get("host_excel")
        if host_excel is None:
            return Response({"error": "请上传 Excel 文件"}, status=status.HTTP_400_BAD_REQUEST)
        if host_excel.size > 2 * 1024 * 1024:
            return Response({"error": "Excel 文件不能超过 2MB"}, status=400)

        buffer = BytesIO()
        for chunk in host_excel.chunks():
            buffer.write(chunk)
        buffer.seek(0)

        try:
            workbook = load_workbook(buffer, read_only=True, data_only=True)
        except Exception:
            return Response({"error": "无效的 Excel 文件"}, status=400)
        worksheet = workbook.worksheets[1] if len(workbook.worksheets) > 1 else workbook.worksheets[0]

        categories = {name: pk for pk, name in HostCategory.objects.values_list("id", "name")}
        strict = getattr(settings, "HOST_SSH_VALIDATE", True)

        created, errors = [], []
        for index, row in enumerate(worksheet.iter_rows(min_row=2), start=2):
            if index > 501:
                errors.append({"row": index, "error": "每次最多导入 500 行"})
                break
            if len(row) < 7:
                errors.append({"row": index, "error": "需要七列数据"})
                continue
            if not row or row[0].value in (None, ""):
                continue

            payload = {
                "category": categories.get(row[0].value),
                "name": row[1].value,
                "ip_addr": row[2].value,
                "port": row[3].value,
                "username": row[4].value,
                "password": "" if row[5].value is None else str(row[5].value),
                "description": row[6].value or "",
            }

            serializer = HostModelSerializer(
                data=payload,
                context={"skip_ssh_validate": not strict, "request": request},
            )
            if serializer.is_valid():
                created.append(serializer.save())
            else:
                errors.append({"row": index, "error": _format_errors(serializer.errors)})

        return Response({
            "data": HostModelSerializer(created, many=True).data,
            "error": errors,
        })


class HostExportView(APIView):
    def get(self, request):
        import csv
        from django.http import HttpResponse
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response['Content-Disposition'] = 'attachment; filename="hosts.csv"'
        response.write('\ufeff')
        writer = csv.writer(response)
        writer.writerow(['名称', '地址', '端口', '用户', '状态', '标签', '备注'])
        for host in accessible_hosts(request.user).iterator():
            cells = [host.name, host.ip_addr, host.port, host.username, host.asset_status, ','.join(host.tags), host.description]
            writer.writerow(["'" + str(cell) if str(cell).startswith(('=', '+', '-', '@', '\t', '\r', '\n')) else cell for cell in cells])
        return response
